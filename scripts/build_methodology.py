import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts"
OUT.mkdir(exist_ok=True)
summary = json.loads((OUT / "analysis-summary.json").read_text(encoding="utf-8"))
expanded = summary["expanded50"]
assignment = summary["assignment15"]
coverage = expanded["descriptive"]
boundaries = summary["evidenceBoundaries"]

BLACK = "000000"
GREEN = "087F5B"
NAVY = "213B52"
PALE = "EFF5F2"
LINE = "D9D9D9"
MUTED = "4F5B54"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.38)
section.bottom_margin = Inches(0.42)
section.left_margin = Inches(0.48)
section.right_margin = Inches(0.48)


def set_font(run, name="Arial", size=9.5, bold=False, color=BLACK):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(2.2)
normal.paragraph_format.line_spacing = 1.0

title_style = doc.styles["Title"]
title_style.font.name = "Georgia"
title_style.font.size = Pt(20)
title_style.font.bold = False
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_after = Pt(2)
title_properties = title_style.element.get_or_add_pPr()
for existing in title_properties.findall(qn("w:pBdr")):
    title_properties.remove(existing)

heading_style = doc.styles["Heading 1"]
heading_style.font.name = "Arial"
heading_style.font.size = Pt(10.5)
heading_style.font.bold = True
heading_style.font.color.rgb = RGBColor.from_string(BLACK)
heading_style.paragraph_format.space_before = Pt(4)
heading_style.paragraph_format.space_after = Pt(2)
heading_style.paragraph_format.keep_with_next = True


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_margins(cell, top=32, start=55, bottom=32, end=55):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def add_borders(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), LINE)
        borders.append(line)
    properties.append(borders)


def format_cell(cell, header=False, fill="FFFFFF", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, NAVY if header else fill)
    add_borders(cell)
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_font(run, size=9.5, bold=header, color="FFFFFF" if header else BLACK)


def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for cell, label, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = label
        format_cell(cell, header=True)
    for row_number, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(value)
            format_cell(cell, fill="FFFFFF" if row_number % 2 == 0 else PALE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.25
    return table


title = doc.add_paragraph("Global Energy Decision Atlas Data and Methodology Note", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.add_heading("Purpose", level=1)
doc.add_paragraph(
    "The atlas screens national markets before local energy diligence. It keeps electricity price, generation structure, primary energy scale, and energy balance as separate measures so a reviewer can see the tradeoffs behind a shortlist."
)

doc.add_heading("Two datasets", level=1)
add_table(
    ["Dataset", "Coverage and periods", "Balance measure"],
    [
        ["Assignment 15", f"{assignment['records']} countries. {assignment['periods']['price']} price, {assignment['periods']['consumption']} consumption and mix, {assignment['periods']['netImports']} trade.", "Reported total energy net imports."],
        ["Expanded 50", f"{expanded['records']} countries. {expanded['periods']['price']} price and {expanded['periods']['consumption']} energy measures.", "Primary energy consumption minus total energy production, used as a screening proxy."],
    ],
    [1.18, 3.75, 2.37],
)

doc.add_heading("Indicators and sources", level=1)
add_table(
    ["Indicator", "Definition and unit", "Primary source"],
    [
        ["Electricity price", "Residential average use or commercial 1,000,000 kWh per year, USD/kWh.", "GlobalPetrolPrices public comparison."],
        ["Non-fossil share", "Solar, wind, hydro, and other share of domestic generation, percent.", "Our World in Data Energy CSV."],
        ["Energy scale", "Total primary energy consumption, EJ.", "Assignment sources or Our World in Data."],
        ["Trade or balance", "Reported net imports for Assignment 15; derived balance gap for Expanded 50, EJ.", "Assignment source or EIA International."],
    ],
    [1.38, 3.8, 2.12],
)

doc.add_heading("Coverage and aggregation", level=1)
doc.add_paragraph(
    f"Expanded 50 covers household price in {coverage['householdPriceUsdPerKwh']['n']}/{coverage['householdPriceUsdPerKwh']['total']} markets and business price in {coverage['businessPriceUsdPerKwh']['n']}/{coverage['businessPriceUsdPerKwh']['total']}. Consumption, electricity mix, and the balance gap each cover {coverage['primaryEnergyConsumptionEj']['n']}/{coverage['primaryEnergyConsumptionEj']['total']}. A fallback may use a value up to two years earlier and remains flagged. Missing values stay N/A."
)
doc.add_paragraph(
    "Regional price uses the unweighted member median. Consumption and balance sum available member values. Regional electricity mix weights member shares by domestic generation. The interface reports overlapping region membership and calculates Pearson r from at least five complete pairs."
)

doc.add_heading("Analysis maturity", level=1)
add_table(
    ["Level", "Status", "What this work supports", "What it does not prove"],
    [
        ["Descriptive", boundaries["descriptive"]["status"], "Coverage, extrema, rankings, concentration, complete-case correlations, and fixed comparisons.", "Causation, future outcomes, or an optimal market."],
        ["Diagnostic", boundaries["diagnostic"]["status"], "Associations and external context can form hypotheses and local questions.", "Why an observed price, mix, demand, reliability, or balance value occurred."],
        ["Predictive", "not performed", "The snapshot establishes a baseline for later forecast design.", "Future electricity prices, demand, reliability, or emissions."],
        ["Prescriptive", "not performed", "A diligence shortlist can order follow-up work.", "A best market, investment return, or optimal site."],
    ],
    [1.0, 1.05, 2.7, 2.55],
)

doc.add_heading("Limits and decision use", level=1)
doc.add_paragraph(
    "The market scope is non-random. Missing values, mixed reference years, flagged fallbacks, national averages, and source-definition differences limit comparison. Correlations use complete cases and have no confidence intervals. The Expanded 50 balance gap is not an observed trade flow. Use the shortlist to plan diligence, not to approve an investment. A site decision needs current industrial tariffs, capital cost, contracts, connection capacity, reliability records, regulatory constraints, and an explicit objective and weighting."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(0)
footer.paragraph_format.space_after = Pt(0)
set_font(footer.add_run("© 2026 Zhekai Li"), size=8.5, color=MUTED)

path = OUT / "global-energy-atlas-methodology.docx"
doc.save(path)
print(path)
